import io
import json

import docx

from app.chunker import chunk_text
from app.services.rag import _clean_history


def test_chunker_respects_size_and_titles():
    text = "# 标题A\n" + "这是一句话。" * 100 + "\n# 标题B\n短内容。"
    chunks = chunk_text(text, chunk_size=100, overlap=20)
    assert all(len(c) <= 100 + len("【标题A】") for c in chunks)
    assert chunks[0].startswith("【标题A】")
    assert chunks[-1] == "【标题B】短内容。"


def test_sample_doc_seeded_and_search(client):
    docs = client.get("/api/documents").json()
    assert len(docs) == 1 and docs[0]["chunk_count"] > 0
    hits = client.post("/api/search", json={"query": "上海出差住宿每晚报销标准", "retrieval": {"top_k": 3}}).json()
    assert len(hits) == 3
    assert hits[0]["doc_name"] == "xinghe_handbook.md"
    assert hits[0]["score"] >= hits[-1]["score"]


def test_document_crud_reindexes(client):
    d = client.post("/api/documents", json={"name": "笔记", "content": "# 猫\n我家的猫叫团子，三岁。"}).json()
    chunks = client.get(f"/api/documents/{d['id']}/chunks").json()
    assert chunks[0]["text"] == "【猫】我家的猫叫团子，三岁。"

    client.put(f"/api/documents/{d['id']}", json={"content": "# 狗\n我家的狗叫可乐。"})
    chunks = client.get(f"/api/documents/{d['id']}/chunks").json()
    assert [c["text"] for c in chunks] == ["【狗】我家的狗叫可乐。"]

    hits = client.post("/api/search", json={"query": "狗叫什么", "doc_ids": [d["id"]]}).json()
    assert hits and all(h["doc_id"] == d["id"] for h in hits)

    assert client.delete(f"/api/documents/{d['id']}").status_code == 204
    assert client.get(f"/api/documents/{d['id']}").status_code == 404
    hits = client.post("/api/search", json={"query": "狗叫什么", "retrieval": {"top_k": 10}}).json()
    assert all(h["doc_id"] != d["id"] for h in hits)


def test_upload_txt_docx_and_reject_unknown(client):
    buf = io.BytesIO()
    w = docx.Document()
    w.add_heading("年会", level=1)
    w.add_paragraph("年会在十二月二十日举行。")
    w.save(buf)
    files = [
        ("files", ("a.txt", "纯文本内容。".encode("gb18030"), "text/plain")),
        ("files", ("b.docx", buf.getvalue(), "application/octet-stream")),
        ("files", ("c.exe", b"xx", "application/octet-stream")),
    ]
    res = client.post("/api/documents/upload", files=files).json()
    assert res[0]["document"]["file_type"] == "txt"
    assert res[1]["document"]["file_type"] == "docx"
    assert "error" in res[2]
    chunks = client.get(f"/api/documents/{res[1]['document']['id']}/chunks").json()
    assert chunks[0]["text"].startswith("【年会】")


def test_settings_and_reindex(client):
    assert client.put("/api/settings", json={"chunk_size": 100, "chunk_overlap": 100}).status_code == 400
    before = client.get("/api/documents").json()[0]["chunk_count"]
    client.put("/api/settings", json={"chunk_size": 100, "chunk_overlap": 20})
    r = client.post("/api/documents/reindex").json()
    after = client.get("/api/documents").json()[0]["chunk_count"]
    assert after > before and r["chunks"] == after


def test_provider_crud_masks_key(client):
    providers = client.get("/api/providers").json()
    assert providers[0]["type"] == "anthropic" and providers[0]["is_default"]
    p = client.post("/api/providers", json={
        "name": "DeepSeek", "type": "openai", "model": "deepseek-chat",
        "base_url": "https://api.deepseek.com", "api_key": "sk-1234567890abcdef",
    }).json()
    assert p["api_key"] == "sk-123…cdef" and p["has_api_key"] and not p["is_default"]
    # 不传 api_key 时保留原 key
    p2 = client.put(f"/api/providers/{p['id']}", json={"model": "deepseek-reasoner"}).json()
    assert p2["model"] == "deepseek-reasoner" and p2["has_api_key"]
    client.post(f"/api/providers/{p['id']}/default")
    assert client.get("/api/settings").json()["default_provider_id"] == p["id"]
    client.delete(f"/api/providers/{p['id']}")
    assert client.get("/api/settings").json()["default_provider_id"] == providers[0]["id"]


def _events(text):
    out = []
    for block in text.strip().split("\n\n"):
        lines = dict(l.split(": ", 1) for l in block.splitlines())
        out.append((lines["event"], json.loads(lines["data"])))
    return out


def test_chat_streams_sources_and_answer(client, monkeypatch):
    seen = {}

    class FakeLLM:
        def __init__(self, cfg):
            seen["cfg"] = cfg

        def stream(self, system, messages):
            seen["messages"] = messages
            yield "答案"
            yield "[1]"

    monkeypatch.setattr("app.services.rag.create_llm", FakeLLM)
    r = client.post("/api/chat", json={
        "query": "年假能留到什么时候？",
        "history": [{"role": "user", "content": "你好"}, {"role": "assistant", "content": "你好！"}],
    })
    ev = _events(r.text)
    assert ev[0][0] == "sources" and len(ev[0][1]) == 4
    assert "".join(d["text"] for e, d in ev if e == "delta") == "答案[1]"
    assert ev[-1][0] == "done"
    assert seen["cfg"].model == "claude-opus-5-5"
    assert [m["role"] for m in seen["messages"]] == ["user", "assistant", "user"]
    assert "<source id=\"1\"" in seen["messages"][-1]["content"]


def test_chat_reports_llm_error(client, monkeypatch):
    from app.llm import LLMError

    class Broken:
        def __init__(self, cfg):
            pass

        def stream(self, system, messages):
            raise LLMError("Claude API Key 无效或未设置")
            yield

    monkeypatch.setattr("app.services.rag.create_llm", Broken)
    ev = _events(client.post("/api/chat", json={"query": "hi"}).text)
    assert ev[-1] == ("error", {"message": "Claude API Key 无效或未设置"})
    ev = _events(client.post("/api/chat", json={"query": "hi", "provider_id": "nope"}).text)
    assert ev[-1][0] == "error"


def test_clean_history():
    h = [{"role": "assistant", "content": "a"}, {"role": "user", "content": "q1"},
         {"role": "assistant", "content": "a1"}, {"role": "user", "content": "q2"}]
    assert _clean_history(h, 3) == [{"role": "user", "content": "q1"}, {"role": "assistant", "content": "a1"}]
    assert _clean_history(h, 0) == []
